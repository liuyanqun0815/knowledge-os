"""Extract plain text from uploaded Office/PDF documents for ingest."""

from __future__ import annotations

from pathlib import Path

from akos.domain.errors import DomainError

TEXT_UPLOAD_SUFFIXES = {".md", ".txt"}
EXTRACTABLE_UPLOAD_SUFFIXES = {".pdf", ".docx"}
ALLOWED_UPLOAD_SUFFIXES = TEXT_UPLOAD_SUFFIXES | EXTRACTABLE_UPLOAD_SUFFIXES

_OCR_ENGINE = None


def _require_docs_extra(package: str) -> None:
    raise DomainError(f"缺少文档解析依赖 {package}；请安装: pip install 'akos[docs]'")


def extract_document(path: Path | str) -> str:
    """Return UTF-8 plain text extracted from path (pdf/docx)."""
    file_path = Path(path)
    if not file_path.is_file():
        raise DomainError(f"file not found: {file_path}")
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        text = _extract_pdf(file_path)
    elif suffix == ".docx":
        text = _extract_docx(file_path)
    else:
        raise DomainError(f"unsupported extract type: {suffix}")
    normalized = _normalize_text(text)
    if not normalized:
        raise DomainError(f"no extractable text in document: {file_path.name}")
    return normalized


def materialize_markdown_for_ingest(path: Path | str) -> Path:
    """Ensure path is ingestible UTF-8 markdown/text; extract binaries to sibling .md."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix in TEXT_UPLOAD_SUFFIXES:
        return file_path
    if suffix not in EXTRACTABLE_UPLOAD_SUFFIXES:
        raise DomainError(f"unsupported file type: {suffix}")
    text = extract_document(file_path)
    md_path = file_path.with_suffix(".md")
    md_path.write_text(text, encoding="utf-8")
    return md_path


def _normalize_text(text: str) -> str:
    lines = [line.rstrip() for line in (text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    collapsed = "\n".join(lines).strip()
    while "\n\n\n" in collapsed:
        collapsed = collapsed.replace("\n\n\n", "\n\n")
    return collapsed


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError:
        _require_docs_extra("python-docx")
    document = Document(str(path))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        value = (paragraph.text or "").strip()
        if value:
            parts.append(value)
    for table in document.tables:
        for row in table.rows:
            cells = [(cell.text or "").strip() for cell in row.cells]
            cells = [cell for cell in cells if cell]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _extract_pdf(path: Path) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        _require_docs_extra("pymupdf")

    document = fitz.open(str(path))
    try:
        parts: list[str] = []
        for page_index in range(len(document)):
            page = document[page_index]
            layer = (page.get_text("text") or "").strip()
            if layer:
                parts.append(layer)
                continue
            ocr_text = _ocr_pdf_page(page)
            if ocr_text:
                parts.append(ocr_text)
        return "\n\n".join(parts)
    finally:
        document.close()


def _ocr_pdf_page(page) -> str:
    try:
        import numpy as np
        from PIL import Image
    except ImportError:
        _require_docs_extra("Pillow/numpy")

    engine = _get_ocr_engine()
    # ~144 dpi equivalent for readable OCR without huge images
    pixmap = page.get_pixmap(matrix=__import__("fitz").Matrix(2, 2), alpha=False)
    image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
    result, _ = engine(np.asarray(image))
    if not result:
        return ""
    return "\n".join(str(item[1]).strip() for item in result if item and item[1])


def _get_ocr_engine():
    global _OCR_ENGINE
    if _OCR_ENGINE is not None:
        return _OCR_ENGINE
    try:
        from rapidocr_onnxruntime import RapidOCR
    except ImportError:
        _require_docs_extra("rapidocr-onnxruntime")
    _OCR_ENGINE = RapidOCR()
    return _OCR_ENGINE
